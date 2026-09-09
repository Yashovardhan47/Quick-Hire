import hashlib
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError


MODEL_VERSION = "secure-document-ingestion-0.3.0"
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}
PROMPT_INJECTION_PATTERNS = (
    "ignore previous instructions",
    "ignore all instructions",
    "system prompt",
    "developer message",
    "assistant:",
    "do not follow the rubric",
)


class DocumentIngestionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    media_type: str
    sha256: str
    size_bytes: int
    page_count: int
    text: str
    segments: list[dict[str, str]]
    security_flags: list[str]


def _safe_filename(filename: str | None) -> str:
    cleaned = PurePath(filename or "resume").name.replace("\x00", "").strip()
    return cleaned[:255] or "resume"


def _clean_text(text: str) -> str:
    text = text.replace("\u200b", "").replace("\ufeff", "")
    text = "".join(character for character in text if character in "\n\t" or ord(character) >= 32)
    return re.sub(r"[ \t]+", " ", text).strip()


def _security_flags(text: str) -> list[str]:
    lowered = text.lower()
    flags = ["instruction_like_content_detected" for phrase in PROMPT_INJECTION_PATTERNS if phrase in lowered]
    if re.search(r"https?://|www\.", lowered):
        flags.append("external_link_present")
    return sorted(set(flags))


def _validate_docx_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2_000:
                raise DocumentIngestionError("DOCX archive contains too many entries")
            if sum(member.file_size for member in members) > 20_000_000:
                raise DocumentIngestionError("DOCX expanded content exceeds the safety limit")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentIngestionError("File is not a valid DOCX document")
            if any(name.startswith(("/", "\\")) or ".." in PurePath(name).parts for name in names):
                raise DocumentIngestionError("DOCX archive contains an unsafe path")
            if any(member.flag_bits & 0x1 for member in members):
                raise DocumentIngestionError("Encrypted DOCX files are not supported")
    except BadZipFile as exc:
        raise DocumentIngestionError("File is not a valid DOCX archive") from exc


def _extract_pdf(data: bytes, max_pages: int) -> tuple[list[dict[str, str]], int, list[str]]:
    if not data.startswith(b"%PDF-"):
        raise DocumentIngestionError("File extension and PDF signature do not match")
    lowered = data.lower()
    if any(marker in lowered for marker in (b"/javascript", b"/openaction", b"/embeddedfiles")):
        raise DocumentIngestionError("PDF contains active or embedded content; upload a flattened document")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise DocumentIngestionError("Encrypted PDF files are not supported")
        if len(reader.pages) > max_pages:
            raise DocumentIngestionError(f"Resume exceeds the {max_pages}-page limit")
        segments = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = _clean_text(page.extract_text() or "")
            if text:
                segments.append({"locator": f"page:{page_number}", "text": text})
        flags = []
        if len(" ".join(segment["text"] for segment in segments)) < max(80, 35 * len(reader.pages)):
            flags.append("low_text_pdf_ocr_required")
        return segments, len(reader.pages), flags
    except DocumentIngestionError:
        raise
    except (PdfReadError, FileNotDecryptedError, KeyError, ValueError) as exc:
        raise DocumentIngestionError("PDF could not be parsed safely") from exc


def _extract_docx(data: bytes) -> tuple[list[dict[str, str]], int]:
    if not data.startswith(b"PK"):
        raise DocumentIngestionError("File extension and DOCX signature do not match")
    _validate_docx_archive(data)
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise DocumentIngestionError("DOCX could not be parsed safely") from exc
    segments = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = _clean_text(paragraph.text)
        if text:
            segments.append({"locator": f"paragraph:{index}", "text": text})
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = _clean_text(" | ".join(cell.text for cell in row.cells))
            if text:
                segments.append({"locator": f"table:{table_index}:row:{row_index}", "text": text})
    return segments, 1


def _extract_txt(data: bytes) -> tuple[list[dict[str, str]], int]:
    if b"\x00" in data[:1_024] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise DocumentIngestionError("Text file contains unsupported binary content")
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentIngestionError("Text file must use UTF-8 or UTF-16 encoding")
    text = _clean_text(text)
    return ([{"locator": "text:1", "text": text}] if text else []), 1


def extract_document(
    filename: str | None,
    data: bytes,
    *,
    max_bytes: int,
    max_pages: int,
    max_characters: int,
) -> ExtractedDocument:
    safe_name = _safe_filename(filename)
    extension = PurePath(safe_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentIngestionError("Supported resume formats are PDF, DOCX and TXT")
    if not data:
        raise DocumentIngestionError("Uploaded resume is empty")
    if len(data) > max_bytes:
        raise DocumentIngestionError(f"Resume exceeds the {max_bytes // 1_048_576} MB size limit")

    if extension == ".pdf":
        segments, page_count, parser_flags = _extract_pdf(data, max_pages)
    elif extension == ".docx":
        segments, page_count = _extract_docx(data)
        parser_flags = []
    else:
        segments, page_count = _extract_txt(data)
        parser_flags = []

    text = "\n".join(segment["text"] for segment in segments).strip()
    if len(text) > max_characters:
        raise DocumentIngestionError("Extracted resume text exceeds the processing limit")
    if len(text) < 40:
        raise DocumentIngestionError("Resume has too little extractable text; use a text-based document")
    flags = sorted(set([*parser_flags, *_security_flags(text)]))
    return ExtractedDocument(
        filename=safe_name,
        media_type=MEDIA_TYPES[extension],
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        page_count=page_count,
        text=text,
        segments=segments,
        security_flags=flags,
    )
